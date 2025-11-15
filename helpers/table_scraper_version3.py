# ============================================================
# 📄 Financial Statement Scraper – MetaLinks Integrated Version
# ============================================================

import json
import re
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font
from typing import Dict, List, Optional, Tuple
import time
import pandas as pd
from IPython.display import display


class FinancialStatementScraper:
    """
    Extracts financial statements from SEC XBRL filings.
    Now integrates MetaLinks.json role detection before pattern matching.
    Includes post-processing year alignment for cash-flow instant shifts.
    """

    def __init__(self, filing_url: str, openai_api_key: str = None):
        self.filing_url = filing_url
        self.openai_api_key = openai_api_key
        self.session = requests.Session()

        self.session.headers.update({
            'User-Agent': 'MyCompany contact@email.com',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
            'Accept-Language': 'en-US,en;q=0.5',
            'Accept-Encoding': 'gzip, deflate, br',
            'Connection': 'keep-alive',
            'Upgrade-Insecure-Requests': '1',
            'Cache-Control': 'max-age=0'
        })

        actual_url = self._extract_document_url(filing_url)
        print("📥 Fetching filing from SEC...")

        for attempt in range(3):
            try:
                time.sleep(0.5)
                resp = self.session.get(actual_url, timeout=30)
                resp.raise_for_status()
                self.html_content = resp.text
                break
            except requests.exceptions.HTTPError as e:
                if e.response.status_code == 403 and attempt < 2:
                    print("⚠ SEC blocked (403). Retrying...")
                    time.sleep((attempt + 1) * 2)
                else:
                    raise Exception(
                        "SEC.gov requires a User-Agent header with contact email. "
                        "Update 'User-Agent' in the code with your own email."
                    )

        self.soup = BeautifulSoup(self.html_content, "html.parser")
        self.tables = self.soup.find_all("table")
        print(f"✓ Loaded HTML with {len(self.tables)} tables")

        self.context_mapping = self._build_context_mapping()
        print(f"✓ Built context mapping with {len(self.context_mapping)} contexts")

        self.metalinks_url = self._construct_metalinks_url(actual_url)
        self.metalinks = self._load_metalinks()

    # ---------------- URL HELPERS ----------------
    def _extract_document_url(self, filing_url: str) -> str:
        if "/ix?doc=" in filing_url:
            return "https://www.sec.gov" + filing_url.split("/ix?doc=")[1]
        return filing_url

    def _construct_metalinks_url(self, document_url: str) -> str:
        return document_url.rsplit("/", 1)[0] + "/MetaLinks.json"

    def _load_metalinks(self) -> Dict:
        try:
            print("📥 Fetching MetaLinks.json...")
            r = self.session.get(self.metalinks_url, timeout=30)
            r.raise_for_status()
            data = r.json()
            if isinstance(data, dict) and "instance" in data:
                first_instance = list(data["instance"].values())[0]
                reports = first_instance.get("report", {})
                print(f"✓ Loaded MetaLinks with {len(reports)} roles")
                return reports
            return {}
        except Exception as e:
            print(f"⚠ Failed to load MetaLinks: {e}")
            return {}

    # ---------------- CONTEXT MAPPING ----------------
    def _build_context_mapping(self) -> Dict[str, Dict[str, str]]:
        mapping = {}
        for ctx in self.soup.find_all(["xbrli:context", "context"]):
            cid = ctx.get("id")
            if not cid:
                continue
            inst = ctx.find(["xbrli:instant", "instant"])
            if inst:
                mapping[cid] = {"date": inst.get_text(strip=True), "type": "instant"}
                continue
            end = ctx.find(["xbrli:enddate", "enddate"])
            start = ctx.find(["xbrli:startdate", "startdate"])
            if end:
                mapping[cid] = {"date": end.get_text(strip=True), "type": "duration"}
            elif start:
                mapping[cid] = {"date": start.get_text(strip=True), "type": "duration"}
        return mapping

    # ---------------- YEAR EXTRACTION ----------------
    def _extract_year_from_context(self, context_ref: str) -> Optional[str]:
        if not context_ref:
            return None
        m = re.search(r"D(\d{4})\d{4}-(\d{4})\d{4}", context_ref)
        if m:
            return m.group(2)
        m = re.search(r"(\d{8})(?!.*\d{8})", context_ref)
        if m:
            return m.group(1)[:4]
        if context_ref in self.context_mapping:
            date = self.context_mapping[context_ref]["date"]
            y = re.search(r"(\d{4})", date)
            if y:
                return y.group(1)
        m = re.search(r"20\d{2}", context_ref)
        return m.group(0) if m else None

    # ---------------- XBRL EXTRACTION + POST-ALIGNMENT ----------------
    def _extract_xbrl_data_from_table(self, table, statement_type: str) -> Tuple[List[str], List[Dict[str, str]]]:
        rows = table.find_all("tr")
        all_years, structured_rows = set(), []

        for row in rows:
            cells = row.find_all(["td", "th"])
            line_item = cells[0].get_text(strip=True) if cells else ""
            year_values = {}

            for cell in cells:
                for tag in cell.find_all(attrs={"contextref": True}):
                    cref = tag.get("contextref", "")
                    year = self._extract_year_from_context(cref)
                    if not year:
                        continue
                    val = tag.get_text(strip=True)
                    meta = {
                        "name": tag.get("name"),
                        "id": tag.get("id"),
                        "unitref": tag.get("unitref"),
                        "decimals": tag.get("decimals"),
                        "format": tag.get("format"),
                        "scale": tag.get("scale"),
                            }
                    
                    year_values[year] = {"value": val, "meta": meta}
                    all_years.add(year)

            if line_item or year_values:
                structured_rows.append({"line_item": line_item, "values": year_values})

        # ========== NOISE FILTER - REMOVE UNWANTED HEADER ROWS ==========
        NOISE_PATTERNS = [
            r'(year|years|month|months|quarter|period)s?\s+(ended|ending)',
            r'^(january|february|march|april|may|june|july|august|september|october|november|december)\s*\d{0,2}',
            r'\(in (millions?|thousands?|billions?|dollars?)\b',
            r'except (per share|share data)',
            r'^\d{4}$|^\d{1,2}/\d{1,2}/\d{2,4}$',
            r'^(as of|for the|fiscal year)',
            r'^\s*$'  # Empty or whitespace only
        ]
        
        structured_rows = [
            r for r in structured_rows 
            if r['values'] or not any(re.search(p, r['line_item'].lower()) for p in NOISE_PATTERNS)
        ]
        # ================================================================

        # dominant year sequence
        year_counts = {}
        for row in structured_rows:
            years = tuple(sorted(row["values"].keys(), reverse=True))
            if len(years) >= 2:
                year_counts[years] = year_counts.get(years, 0) + 1
        dominant_years = []
        if year_counts:
            dominant_years = max(year_counts, key=year_counts.get)

        # shift lagging instantaneous rows (cash flow only)
        if statement_type == "cash_flow" and dominant_years:
            dominant_years_int = [int(y) for y in dominant_years]
            for row in structured_rows:
                current_years = sorted([int(y) for y in row["values"].keys()], reverse=True)
                if len(current_years) == len(dominant_years) and all(
                    (dy - cy == 1) for dy, cy in zip(dominant_years_int, current_years)
                ):
                    shifted = {str(int(y) + 1): v for y, v in row["values"].items()}
                    row["values"] = shifted

        all_years = set()
        for r in structured_rows:
            all_years.update(r["values"].keys())

        return sorted(all_years, reverse=True), structured_rows

    # ---------------- TABLE / EXCEL HELPERS ----------------
    def extract_table_data(self, table_idx: int, statement_type: str) -> List[List[str]]:
        if table_idx >= len(self.tables):
            return []
        table = self.tables[table_idx]
        xbrl_tags = table.find_all(attrs={"contextref": True})
        if xbrl_tags:
            years, rows = self._extract_xbrl_data_from_table(table, statement_type)
            if years and rows:
                data = [["Line Item"] + years]
                for r in rows:
                    row = [r["line_item"]] + [r["values"].get(y, "") for y in years]
                    if any(row):
                        data.append(row)
                print(f"✓ XBRL extraction: {len(years)} columns, {len(data)-1} rows")
                return data
        return self._extract_table_data_traditional(table)

    def _extract_table_data_traditional(self, table) -> List[List[str]]:
        try:
            from io import StringIO
            import warnings
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                dfs = pd.read_html(StringIO(str(table)), flavor="html5lib")
            if dfs:
                df = dfs[0]
                data = [df.columns.tolist()] + df.values.tolist()
                cleaned = [[str(c).strip() if pd.notna(c) else "" for c in row] for row in data if any(row)]
                return cleaned
        except Exception:
            pass
        rows = []
        for tr in table.find_all("tr"):
            row = [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]
            if any(row):
                rows.append(row)
        return rows

    # ---------------- STATEMENT EXTRACTION ----------------
    def extract_statement(self, role_id: str, statement_name: str, statement_type: str,
                          output_filename: str, display_output: bool = True) -> Dict:
        print(f"\n{'='*80}\nExtracting: {statement_name}\n{'='*80}\n")

        anchor_idx = self.find_table_by_unique_anchor(role_id, statement_type)
        if anchor_idx is None:
            print("⚠ Using pattern matching fallback...")
            kw = {
                "cash_flow": ["Cash flows", "Operating activities", "Investing activities", "Financing activities"],
                "balance_sheet": ["Assets", "Liabilities", "Cash and cash equivalents"],
                "income_statement": ["Revenues", "Net earnings", "Operating", "Income"],
            }.get(statement_type, [])
            matches = self.find_table_by_pattern(kw)
            if matches:
                anchor_idx = matches[0]
            else:
                return {"status": "failed", "error": f"Could not locate {statement_name}"}
        
        data = self.extract_table_data(anchor_idx, statement_type)
        
        # If extract_table_data returns XBRL-structured rows
        if isinstance(data, dict) and "rows" in data:
            json_output = data
        else:
            # Convert tabular data → JSON-like format
            if not data or len(data) < 2:
                return {"status": "failed", "error": f"No data found for {statement_name}"}
            years = data[0][1:]
            rows = []
            for r in data[1:]:
                label = r[0]
                vals = {y: v for y, v in zip(years, r[1:]) if v != ""}
                rows.append({"line_item": label, "values": vals})
            json_output = {"statement_type": statement_type, "years": years, "rows": rows}
        
        # --- Save clean JSON ---
        json_path = output_filename.replace(".xlsx", ".json")
        with open(json_path, "w") as f:
            json.dump(json_output, f, indent=2)
        
        print(f"✅ Exported clean JSON → {json_path}")
        return {"status": "success", "output_file": json_path, "json": json_output}


    def extract_all_statements(self, display_output: bool = True) -> Dict[str, Dict]:
        configs = {
            "balance_sheet": ("balance_sheet", "Consolidated_Balance_Sheets"),
            "income_statement": ("income_statement", "Consolidated_Statements_of_Earnings"),
            "cash_flow": ("cash_flow", "Consolidated_Statements_of_Cash_Flows"),
        }
        results = {}
        for key, (stype, fname) in configs.items():
            results[key] = self.extract_statement(None, fname.replace("_", " "), stype, f"{fname}.xlsx", display_output)
        return results

    # ---------------- UTILITIES ----------------
    def to_dataframe(self, data: List[List[str]]) -> pd.DataFrame:
        if not data or len(data) < 2:
            return pd.DataFrame()
        df = pd.DataFrame(data[1:], columns=data[0])
        df = df[df.apply(lambda x: x.astype(str).str.strip().ne("").any(), axis=1)]
        return df

    def save_to_excel(self, data: List[List[str]], sheet_name: str, output_path: str):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]
        bold = Font(bold=True)
        for i, row in enumerate(data, 1):
            for j, val in enumerate(row, 1):
                c = ws.cell(i, j, val)
                if i == 1:
                    c.font = bold
        wb.save(output_path)

    # ---------------- META-LINK TABLE MATCHING ----------------
    def find_table_by_unique_anchor(self, role_id: Optional[str], statement_type: str) -> Optional[int]:
        """
        Locate table by using MetaLinks.json uniqueAnchor before pattern search.
        Falls back to None if not found.
        """
        if not self.metalinks:
            return None

        TAXONOMY_MAP = {
            "balance_sheet": [
                "consolidated balance sheets",
                "balance sheet",
                "statement of financial position",
                "financial condition",
                "assets and liabilities",
            ],
            "income_statement": [
                "consolidated statements of operations",
                "consolidated statement of operations",
                "income statement",
                "income statements",
                "consolidated statements of profit or loss",
                "statement of earnings",
                "profit and loss",
                "consolidated income statements",
                "consolidated statements of income",
                "consolidated statements of earnings",
            ],
            "cash_flow": [
                "consolidated statements of cash flows",
                "consolidated statement of cash flows",
                "statement of cash flows",
                "cash flows statements"
            ],
        }

        statement_roles = {
            rid: r for rid, r in self.metalinks.items()
            if r.get("groupType", "").lower() == "statement"
        }

        # Build lookup by shortname
        role_lookup = {}
        for rid, rpt in statement_roles.items():
            shortname = rpt.get("shortName", "").lower().strip()
            for stype, names in TAXONOMY_MAP.items():
                if any(shortname == n.lower() for n in names):
                    role_lookup[stype] = (rid, rpt)
                    break

        # 1️⃣ if explicit role id given
        if role_id and role_id in statement_roles:
            role = statement_roles[role_id]
            anchor = role.get("uniqueAnchor", {}).get("name")
            if anchor:
                for idx, tbl in enumerate(self.tables):
                    if tbl.find(attrs={"name": anchor}):
                        print(f"✓ MetaLinks direct match for {role_id} → table {idx}")
                        return idx

        # 2️⃣ else detect by statement_type taxonomy match
        if statement_type in role_lookup:
            _, role = role_lookup[statement_type]
            anchor_name = role.get("uniqueAnchor", {}).get("name")
            if anchor_name:
                for idx, tbl in enumerate(self.tables):
                    if tbl.find(attrs={"name": anchor_name}):
                        print(f"✓ MetaLinks matched {statement_type} → table {idx}")
                        return idx

        return None

    def find_table_by_pattern(self, keywords: List[str], min_length: int = 800) -> List[int]:
        found = []
        for i, t in enumerate(self.tables):
            text = t.get_text().lower()
            if len(text) >= min_length and all(k.lower() in text for k in keywords):
                found.append(i)
        return found


print("✓ Financial Statement Scraper loaded successfully!")
print("✓ Integrated MetaLinks-based statement detection")
print("✓ Pattern-based fallback retained")
print("✓ Post-alignment year correction active")
print("✓ Noise filter integrated for clean data extraction\n")